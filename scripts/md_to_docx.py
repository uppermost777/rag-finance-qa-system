"""把 Markdown 论文初稿转换为 Word 文档（供毕业设计论文排版直接使用）。

支持的语法子集：一级/二级/三级标题、正文段落、无序列表、引用块（>）、代码块（```）、
Markdown 表格（转为 Word 表格）、加粗（**...**）与行内代码（`` `...` ``）。
公式与 Mermaid 图请在 Word 中自行插入图片（脚本会把它们保留为代码块文本）。

用法：
    python scripts/md_to_docx.py --input ../outputs/论文-第3-4章-系统设计与实现.md \
        --output ../outputs/论文初稿-第3-4章.docx --title "第3-4章"
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt  # noqa: E402

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_INLINE_CODE = re.compile(r"`([^`]+)`")
_HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
_TABLE_ROW = re.compile(r"^\|(.+)\|$")
_SEPARATOR_ROW = re.compile(r"^\|[\s:\-|]+\|$")


def _add_runs(paragraph, text: str) -> None:  # noqa: ANN001 - python-docx 对象
    """写入文本，处理 **加粗** 与 `行内代码`。"""
    text = _INLINE_CODE.sub(r"\1", text)
    position = 0
    for match in _BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _flush_table(document: Document, rows: list[list[str]]) -> None:  # noqa: ANN001
    """把收集到的 Markdown 表格写成 Word 表格。"""
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    for row_values in rows:
        cells = table.add_row().cells
        for index in range(columns):
            value = row_values[index] if index < len(row_values) else ""
            cells[index].text = _INLINE_CODE.sub(r"\1", value).replace("**", "")


def convert(markdown_path: Path, output_path: Path, title: str | None = None) -> Path:
    """执行转换。"""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)

    if title:
        document.add_heading(title, level=0)

    table_rows: list[list[str]] = []
    in_code_block = False
    for raw_line in markdown_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()

        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block:
            paragraph = document.add_paragraph(line)
            paragraph.style = document.styles["No Spacing"]
            if paragraph.runs:  # 空行不会产生 run
                paragraph.runs[0].font.name = "Consolas"
                paragraph.runs[0].font.size = Pt(9)
            continue

        if _TABLE_ROW.match(line):
            if _SEPARATOR_ROW.match(line):
                continue
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            table_rows.append(cells)
            continue
        if table_rows:
            _flush_table(document, table_rows)
            table_rows = []

        heading = _HEADING.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_heading(heading.group(2).strip(), level=level)
            continue

        if line.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            _add_runs(paragraph, line.lstrip("> ").strip())
            for run in paragraph.runs:
                run.italic = True
            continue

        if re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, re.sub(r"^[-*]\s+", "", line))
            continue

        if re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            _add_runs(paragraph, re.sub(r"^\d+\.\s+", "", line))
            continue

        if not line.strip():
            continue

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Pt(21)  # 首行缩进两字符
        _add_runs(paragraph, line)

    if table_rows:
        _flush_table(document, table_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def main() -> int:
    """脚本入口。"""
    parser = argparse.ArgumentParser(description="Markdown 转 Word（论文初稿）")
    parser.add_argument("--input", required=True, help="Markdown 文件路径")
    parser.add_argument("--output", required=True, help="输出 .docx 路径")
    parser.add_argument("--title", default=None, help="文档标题（可选）")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.is_file():
        print(f"输入文件不存在：{input_path}")
        return 1
    result = convert(input_path, Path(args.output), args.title)
    print(f"已生成：{result}（{result.stat().st_size / 1024:.1f} KB）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
